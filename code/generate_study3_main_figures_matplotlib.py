from __future__ import annotations

import json
import math
import re
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.patches import FancyArrowPatch, Rectangle
from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
INPUT_XLSX = ROOT / "Table" / "Study3_All_Tables.xlsx"
OUT_DIR = ROOT / "Figure" / "Study3_main"
FIGURE1_STEM = "Figure_1_subtype_descriptive_outcomes"
LEGACY_FIGURE1_STEM = "Figure_1_subtype_landscape"

WIDTH_IN = 7.2
HEIGHT_IN = 5.0
PREVIEW_DPI = 250
TIFF_DPI = 600
PANEL_TOP = 0.92

PALETTE = {
    "C1": "#00A1D5",
    "C2": "#DF8F44",
    "C3": "#B24745",
    "Before mediators": "#6F7782",
    "After mediators": "#374E55",
    "ink": "#374E55",
    "muted": "#6F6F6F",
    "grid": "#E5E8EA",
    "zero": "#A8B0B7",
    "light": "#F5F7F8",
    "line": "#374E55",
}

SUBTYPES = [
    ("C1", "Low-stable"),
    ("C2", "Low-mod / high-fluct"),
    ("C3", "High / high-fluct"),
]

MEDIATORS = [
    "Entrapment",
    "Perceived burdensomeness",
    "Thwarted belongingness",
]

PROFILE_LABELS = {
    "Stress sum intensity": "Stress",
    "Entrapment": "Entrapment",
    "Perceived burdensomeness": "Burdensomeness",
    "Thwarted belongingness": "Belongingness",
    "Suicidal ideation": "SI",
}

MEDIATOR_LABELS = {
    "Entrapment": "Entrapment",
    "Perceived burdensomeness": "Burdensomeness",
    "Thwarted belongingness": "Belongingness",
}


def configure_style() -> None:
    mpl.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans", "sans-serif"],
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
            "font.size": 7,
            "axes.labelsize": 7,
            "axes.titlesize": 8,
            "xtick.labelsize": 6.4,
            "ytick.labelsize": 6.8,
            "axes.linewidth": 0.55,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "legend.frameon": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )


def rows_for_sheet(wb, sheet: str) -> list[list]:
    ws = wb[sheet]
    rows: list[list] = []
    for row in ws.iter_rows(values_only=True):
        if any(v is not None for v in row):
            rows.append(list(row))
    return rows


def normalize_minus(value: str) -> str:
    return value.replace("\u2212", "-").replace("\u2013", "-").replace("\u2014", "-")


def parse_float(value) -> float:
    return float(str(value).replace(",", "").strip())


EST_RE = re.compile(
    r"^\s*([+-]?\d+(?:\.\d+)?)\s*\[\s*([+-]?\d+(?:\.\d+)?)\s*,\s*([+-]?\d+(?:\.\d+)?)\s*\]\s*(\*)?\s*$"
)


def parse_est_ci(value: str) -> dict[str, float | bool]:
    text = normalize_minus(str(value).strip())
    match = EST_RE.match(text)
    if not match:
        raise ValueError(f"Cannot parse estimate interval: {value!r}")
    est, low, high, star = match.groups()
    low_f = float(low)
    high_f = float(high)
    return {
        "estimate": float(est),
        "low": low_f,
        "high": high_f,
        "credible": bool(star) or low_f > 0 or high_f < 0,
    }


def parse_mean_sd(value: str) -> tuple[float, float]:
    match = re.match(r"^([+-]?\d+(?:\.\d+)?)\s*\(([+-]?\d+(?:\.\d+)?)\)$", str(value).replace(",", "").strip())
    if not match:
        raise ValueError(f"Cannot parse mean SD: {value!r}")
    return float(match.group(1)), float(match.group(2))


def parse_n_percent(value: str) -> tuple[int, float | None]:
    match = re.match(r"^(\d+)(?:\s*\(([+-]?\d+(?:\.\d+)?)\))?$", str(value).replace(",", "").strip())
    if not match:
        raise ValueError(f"Cannot parse n percent: {value!r}")
    return int(match.group(1)), float(match.group(2)) if match.group(2) else None


def extract_data() -> dict[str, pd.DataFrame]:
    wb = load_workbook(INPUT_XLSX, data_only=True)
    data: dict[str, pd.DataFrame] = {}

    t1 = rows_for_sheet(wb, "Table 1")
    by_name = {r[0]: r for r in t1 if r and r[0]}
    sample = pd.DataFrame(
        [
            {"metric": "Participants", "label": "Participants with >=1 person-wave", "value": parse_float(by_name["Unique participants contributing ≥1 person-wave to subtype, n"][1])},
            {"metric": "Person-waves", "label": "Person-waves", "value": parse_float(by_name["Person-waves, n (% of total)"][1])},
            {"metric": "EMA prompts", "label": "EMA prompts", "value": parse_float(by_name["EMA prompts, n"][1])},
            {"metric": "EMA waves", "label": "EMA waves", "value": 4},
        ]
    )
    data["sample"] = sample

    total_prompts = parse_float(by_name["EMA prompts, n"][1])
    comp_rows = []
    for i, (code, label) in enumerate(SUBTYPES, start=2):
        n, pct = parse_n_percent(by_name["Person-waves, n (% of total)"][i])
        comp_rows.append({"measure": "Person-waves", "subtype": code, "subtype_label": label, "n": n, "percent": pct})
        prompts = parse_float(by_name["EMA prompts, n"][i])
        comp_rows.append({"measure": "EMA prompts", "subtype": code, "subtype_label": label, "n": int(prompts), "percent": prompts / total_prompts * 100})
    data["composition"] = pd.DataFrame(comp_rows)

    profile_rows = []
    for variable in [
        "Stress sum intensity",
        "Entrapment",
        "Perceived burdensomeness",
        "Thwarted belongingness",
        "Suicidal ideation",
    ]:
        row = by_name[variable]
        for i, (code, label) in enumerate(SUBTYPES, start=2):
            mean, sd = parse_mean_sd(row[i])
            profile_rows.append({"variable": variable, "variable_label": PROFILE_LABELS[variable], "subtype": code, "subtype_label": label, "mean": mean, "sd": sd})
    profile = pd.DataFrame(profile_rows)
    profile["standardized_mean"] = profile.groupby("variable")["mean"].transform(
        lambda values: (values - values.mean()) / values.std(ddof=0) if values.std(ddof=0) else np.repeat(0.0, len(values))
    )
    data["profile"] = profile

    prevalence_rows = []
    for variable in ["Any momentary stressor endorsed", "Any momentary suicidal ideation (SI > 1)"]:
        row = by_name[variable]
        for i, (code, label) in enumerate(SUBTYPES, start=2):
            prevalence_rows.append({"variable": variable, "subtype": code, "subtype_label": label, "percent": parse_float(row[i])})
    data["prevalence"] = pd.DataFrame(prevalence_rows)

    t2 = rows_for_sheet(wb, "Table 2")
    section = None
    overall_rows = []
    for row in t2:
        first = row[0]
        if first in [
            "a-paths: Stress → Mediator",
            "b-paths: Mediator → Suicidal ideation",
            "Within-person indirect effects (a × b)",
            "Dominance contrasts among indirect effects",
        ]:
            section = first
            continue
        if section and row[1] is not None and "[" in str(row[2]):
            if section.startswith("a-paths"):
                panel = "a-path"
            elif section.startswith("b-paths"):
                panel = "b-path"
            elif section.startswith("Within"):
                panel = "indirect"
            else:
                panel = "dominance"
            mediator = next((m for m in MEDIATORS if m in str(first)), str(first))
            overall_rows.append({"panel": panel, "path": first, "mediator": mediator, "pd": row[3], **parse_est_ci(f"{row[1]} {row[2]}")})
    data["overall_mediation"] = pd.DataFrame(overall_rows)

    t3 = rows_for_sheet(wb, "Table 3")
    subtype_rows = []
    contrast_rows = []
    current_mediator = None
    for row in t3:
        first = row[0]
        if first in MEDIATORS:
            current_mediator = first
            continue
        if not current_mediator or first is None:
            continue
        first_s = str(first)
        if first_s.startswith("C"):
            subtype = first_s.split()[0]
            for panel, value in [("a-path", row[1]), ("b-path", row[2]), ("indirect", row[3])]:
                subtype_rows.append({"panel": panel, "mediator": current_mediator, "subtype": subtype, **parse_est_ci(value)})
        elif first_s.startswith("Indirect contrast"):
            contrast_rows.append(
                {
                    "panel": "indirect_contrast",
                    "mediator": current_mediator,
                    "contrast": normalize_minus(first_s.replace("Indirect contrast: ", "")),
                    **parse_est_ci(str(row[3]).replace("*", " *")),
                }
            )
    data["subtype_mediation"] = pd.DataFrame(subtype_rows)
    data["subtype_contrasts"] = pd.DataFrame(contrast_rows)

    t4 = rows_for_sheet(wb, "Table 4")
    within_rows = []
    contrast4_rows = []
    section = None
    for row in t4:
        first = row[0]
        if first == "Within-class stress → SI slope":
            section = "within"
            continue
        if first == "Subtype contrasts on stress → SI slope":
            section = "contrast"
            continue
        if section and first and row[1] is not None and "[" in str(row[1]):
            label = str(first).split()[0] if section == "within" else normalize_minus(str(first))
            dest = within_rows if section == "within" else contrast4_rows
            for model, value in [("Before mediators", row[1]), ("After mediators", row[2]), ("Attenuation", row[3])]:
                dest.append({"panel": section, "label": label, "model": model, **parse_est_ci(value)})
    data["attenuation_within"] = pd.DataFrame(within_rows)
    data["attenuation_contrasts"] = pd.DataFrame(contrast4_rows)
    return data


def save_source_data(data: dict[str, pd.DataFrame]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    profile_source = data["profile"][["variable", "variable_label", "subtype", "subtype_label", "mean", "sd"]].copy()
    profile_source.insert(0, "section", "raw_mean")
    prevalence_source = data["prevalence"][["variable", "subtype", "subtype_label", "percent"]].copy()
    prevalence_source["variable_label"] = prevalence_source["variable"].replace(
        {
            "Any momentary stressor endorsed": "Any stressor",
            "Any momentary suicidal ideation (SI > 1)": "Any SI",
        }
    )
    prevalence_source["mean"] = np.nan
    prevalence_source["sd"] = np.nan
    prevalence_source.insert(0, "section", "prevalence")
    figure1_source = pd.concat(
        [
            profile_source,
            prevalence_source[["section", "variable", "variable_label", "subtype", "subtype_label", "mean", "sd", "percent"]],
        ],
        ignore_index=True,
        sort=False,
    )
    figure1_source.to_csv(OUT_DIR / f"{FIGURE1_STEM}_source_data.csv", index=False, encoding="utf-8-sig")
    data["overall_mediation"].to_csv(OUT_DIR / "Figure_2_overall_mediation_source_data.csv", index=False, encoding="utf-8-sig")
    pd.concat([data["subtype_mediation"], data["subtype_contrasts"]], ignore_index=True, sort=False).to_csv(
        OUT_DIR / "Figure_3_subtype_specific_mediation_source_data.csv", index=False, encoding="utf-8-sig"
    )
    data["subtype_mediation"].to_csv(OUT_DIR / "Figure_S1_path_decomposition_source_data.csv", index=False, encoding="utf-8-sig")
    pd.concat([data["attenuation_within"], data["attenuation_contrasts"]], ignore_index=True, sort=False).to_csv(
        OUT_DIR / "Figure_4_stress_si_attenuation_source_data.csv", index=False, encoding="utf-8-sig"
    )


def cleanup_legacy_outputs() -> None:
    for ext in [".svg", ".pdf", ".png", ".tiff"]:
        path = OUT_DIR / f"{LEGACY_FIGURE1_STEM}{ext}"
        if path.exists():
            path.unlink()
    legacy_source = OUT_DIR / f"{LEGACY_FIGURE1_STEM}_source_data.csv"
    if legacy_source.exists():
        legacy_source.unlink()


def save_figure(fig: mpl.figure.Figure, stem: str) -> None:
    base = OUT_DIR / stem
    fig.savefig(base.with_suffix(".svg"))
    fig.savefig(base.with_suffix(".pdf"))
    fig.savefig(base.with_suffix(".png"), dpi=PREVIEW_DPI)
    fig.savefig(base.with_suffix(".tiff"), dpi=TIFF_DPI)
    plt.close(fig)


def panel_label(ax, label: str, x: float = -0.12, y: float = 1.08) -> None:
    ax.text(x, y, label, transform=ax.transAxes, fontsize=10.5, fontweight="bold", va="top", ha="left", color=PALETTE["ink"])


def subtype_color(label: str) -> str:
    if label.startswith("C2"):
        return PALETTE["C2"]
    if label.startswith("C3"):
        return PALETTE["C3"]
    return PALETTE["C1"]


def clean_axis(ax) -> None:
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(length=2.8, width=0.55, colors=PALETTE["muted"])
    ax.xaxis.label.set_color(PALETTE["muted"])
    ax.yaxis.label.set_color(PALETTE["muted"])


def draw_interval(ax, df: pd.DataFrame, label_col: str, color_col: str | None, xmin: float, xmax: float, xticks: list[float], xlabel: str, zero: bool = False) -> None:
    y = np.arange(len(df))[::-1]
    ax.set_xlim(xmin, xmax)
    ax.set_yticks(y)
    ax.set_yticklabels(df[label_col], color=PALETTE["ink"])
    ax.set_xticks(xticks)
    if zero:
        ax.axvline(0, color=PALETTE["zero"], lw=0.65, zorder=0)
    ax.grid(axis="y", color=PALETTE["grid"], lw=0.55)
    for yy, (_, row) in zip(y, df.iterrows()):
        color = PALETTE.get(str(row[color_col]), PALETTE["ink"]) if color_col else row.get("color", PALETTE["ink"])
        ax.errorbar(
            row["estimate"],
            yy,
            xerr=[[row["estimate"] - row["low"]], [row["high"] - row["estimate"]]],
            fmt="o",
            color=color,
            ecolor=color,
            elinewidth=1.4,
            capsize=3,
            markersize=4.4,
            markeredgecolor="white",
            markeredgewidth=0.5,
            zorder=3,
        )
    ax.set_xlabel(xlabel)
    clean_axis(ax)


def figure1(data: dict[str, pd.DataFrame]) -> None:
    fig = plt.figure(figsize=(WIDTH_IN, HEIGHT_IN), constrained_layout=False)
    outer = fig.add_gridspec(2, 1, height_ratios=[1.45, 0.9])
    fig.subplots_adjust(left=0.075, right=0.98, bottom=0.11, top=0.91, hspace=0.58)
    means_grid = outer[0].subgridspec(2, 3, wspace=0.42, hspace=0.55)
    mean_axes = [fig.add_subplot(means_grid[i // 3, i % 3]) for i in range(5)]
    legend_ax = fig.add_subplot(means_grid[1, 2])
    ax_b = fig.add_subplot(outer[1])

    vars_order = ["Stress sum intensity", "Entrapment", "Perceived burdensomeness", "Thwarted belongingness", "Suicidal ideation"]
    x = np.arange(len(SUBTYPES))
    bar_width = 0.66
    for idx, (ax, variable) in enumerate(zip(mean_axes, vars_order)):
        subset = data["profile"][data["profile"]["variable"] == variable].copy()
        subset["order"] = subset["subtype"].map({"C1": 1, "C2": 2, "C3": 3})
        subset = subset.sort_values("order")
        colors = [PALETTE[code] for code in subset["subtype"]]
        ax.bar(
            x,
            subset["mean"],
            yerr=subset["sd"],
            width=bar_width,
            color=colors,
            edgecolor="white",
            linewidth=0.7,
            error_kw={"elinewidth": 0.8, "ecolor": PALETTE["ink"], "capsize": 2.5, "capthick": 0.8},
        )
        ymin = min(0.0, float((subset["mean"] - subset["sd"]).min()) * 1.08)
        ymax = float((subset["mean"] + subset["sd"]).max()) * 1.14
        ax.set_ylim(ymin, ymax)
        ax.set_title(PROFILE_LABELS[variable], loc="left", pad=6, fontweight="bold")
        ax.set_xticks(x)
        ax.set_xticklabels([code for code, _ in SUBTYPES])
        if idx in [0, 3]:
            ax.set_ylabel("Mean (SD)")
        ax.grid(axis="y", color=PALETTE["grid"], lw=0.55)
        clean_axis(ax)
        if idx == 0:
            panel_label(ax, "A", -0.18, 1.18)
            ax.text(-0.18, 1.30, "Raw momentary means", transform=ax.transAxes, fontsize=9.2, fontweight="bold", color=PALETTE["ink"])

    legend_ax.axis("off")
    handles = [mpl.patches.Patch(facecolor=PALETTE[c], edgecolor="white", label=f"{c}: {l.replace(' / ', '/')}") for c, l in SUBTYPES]
    legend_ax.legend(handles=handles, loc="center left", frameon=False, fontsize=6.8, handlelength=1.0, handleheight=0.9, borderaxespad=0)

    panel_label(ax_b, "B", -0.065, 1.16)
    ax_b.set_title("Prompt-level prevalence", loc="left", pad=8, fontweight="bold")
    prevalence_order = [
        ("Any momentary stressor endorsed", "Any stressor"),
        ("Any momentary suicidal ideation (SI > 1)", "Any SI"),
    ]
    group_x = np.arange(len(prevalence_order))
    offsets = np.array([-0.24, 0.0, 0.24])
    width = 0.22
    for subtype_index, (code, _) in enumerate(SUBTYPES):
        values = []
        for variable, _ in prevalence_order:
            row = data["prevalence"][(data["prevalence"]["variable"] == variable) & (data["prevalence"]["subtype"] == code)].iloc[0]
            values.append(row["percent"])
        xpos = group_x + offsets[subtype_index]
        ax_b.bar(xpos, values, width=width, color=PALETTE[code], edgecolor="white", linewidth=0.7)
        for xx, value in zip(xpos, values):
            ax_b.text(xx, value + 2.2, f"{value:.1f}", ha="center", va="bottom", fontsize=5.8, color=PALETTE[code], fontweight="bold")
    ax_b.set_xlim(-0.55, len(prevalence_order) - 0.45)
    ax_b.set_ylim(0, 105)
    ax_b.set_xticks(group_x)
    ax_b.set_xticklabels([label for _, label in prevalence_order], fontweight="bold")
    ax_b.set_ylabel("Prompt-level prevalence (%)")
    ax_b.set_yticks([0, 25, 50, 75, 100])
    ax_b.grid(axis="y", color=PALETTE["grid"], lw=0.55)
    clean_axis(ax_b)
    save_figure(fig, FIGURE1_STEM)


def figure2(data: dict[str, pd.DataFrame]) -> None:
    fig = plt.figure(figsize=(WIDTH_IN, HEIGHT_IN), constrained_layout=False)
    gs = fig.add_gridspec(3, 2, width_ratios=[1.15, 1], height_ratios=[1.05, 0.9, 0.95])
    fig.subplots_adjust(left=0.13, right=0.98, bottom=0.09, top=PANEL_TOP, wspace=0.72, hspace=0.78)
    ax_a = fig.add_subplot(gs[:2, 0])
    ax_d = fig.add_subplot(gs[2, 0])
    ax_b = fig.add_subplot(gs[0, 1])
    ax_c = fig.add_subplot(gs[1, 1])
    ax_e = fig.add_subplot(gs[2, 1])

    ax_a.axis("off")
    panel_label(ax_a, "A", -0.04, 1.02)
    ax_a.set_title("Evidence logic", loc="left", pad=8, fontweight="bold")
    ax_a.text(0.10, 0.48, "Stress", ha="center", va="center", fontsize=9, fontweight="bold", bbox=dict(boxstyle="square,pad=0.55", fc=PALETTE["light"], ec="#D9DEE5"))
    ax_a.text(0.90, 0.48, "SI", ha="center", va="center", fontsize=9, fontweight="bold", bbox=dict(boxstyle="square,pad=0.55", fc=PALETTE["light"], ec="#D9DEE5"))
    med_pos = [(0.50, 0.75, "Entrapment"), (0.50, 0.48, "Perceived burdensomeness"), (0.50, 0.21, "Thwarted belongingness")]
    for x, y, med in med_pos:
        label = med.replace("Perceived burdensomeness", "Burdensomeness").replace("Thwarted belongingness", "Belongingness")
        ax_a.text(x, y, label, ha="center", va="center", fontsize=7, fontweight="bold", color=PALETTE["ink"], bbox=dict(boxstyle="square,pad=0.45", fc="white", ec=PALETTE["line"], lw=1.0))
        ax_a.add_patch(FancyArrowPatch((0.18, 0.48), (x - 0.13, y), transform=ax_a.transAxes, arrowstyle="-|>", mutation_scale=10, color=PALETTE["line"], lw=1.0))
        ax_a.add_patch(FancyArrowPatch((x + 0.13, y), (0.82, 0.48), transform=ax_a.transAxes, arrowstyle="-|>", mutation_scale=10, color=PALETTE["line"], lw=1.0))
    ax_a.text(0.5, 0.04, "All three indirect pathways were credibly positive.", ha="center", fontsize=8, fontweight="bold")

    overall = data["overall_mediation"]
    for ax, panel, title_text, xlim, ticks in [
        (ax_b, "a-path", "a-paths: stress to mediator", (0.03, 0.08), [0.03, 0.05, 0.07]),
        (ax_c, "b-path", "b-paths: mediator to SI", (0.0, 0.065), [0, 0.02, 0.04, 0.06]),
        (ax_d, "indirect", "Indirect effects", (0, 0.0035), [0, 0.001, 0.002, 0.003]),
    ]:
        rows = overall[overall["panel"] == panel].copy()
        rows["label"] = rows["mediator"].replace({"Perceived burdensomeness": "Burdensomeness", "Thwarted belongingness": "Belongingness"})
        rows["color"] = PALETTE["line"]
        panel_label(ax, {"a-path": "B", "b-path": "C", "indirect": "D"}[panel], -0.14, 1.13)
        ax.set_title(title_text, loc="left", pad=8, fontweight="bold")
        draw_interval(ax, rows, "label", None, xlim[0], xlim[1], ticks, "Posterior median" if panel != "indirect" else "Indirect effect (a x b)")

    dom = overall[overall["panel"] == "dominance"].copy()
    dom["label"] = dom["path"].str.replace("Perceived burdensomeness", "Burdensomeness", regex=False).str.replace("Thwarted belongingness", "Belongingness", regex=False).str.replace(" - ", "\n- ", regex=False)
    panel_label(ax_e, "E", -0.14, 1.13)
    ax_e.set_title("Dominance contrasts", loc="left", pad=8, fontweight="bold")
    dom["color"] = PALETTE["line"]
    draw_interval(ax_e, dom, "label", None, -0.0006, 0.0026, [-0.0005, 0, 0.001, 0.002], "Difference in indirect effect", zero=True)
    save_figure(fig, "Figure_2_overall_mediation")


def figure3(data: dict[str, pd.DataFrame]) -> None:
    fig = plt.figure(figsize=(WIDTH_IN, HEIGHT_IN), constrained_layout=False)
    gs = fig.add_gridspec(3, 2, width_ratios=[1.05, 1], height_ratios=[1, 1, 1])
    fig.subplots_adjust(left=0.13, right=0.98, bottom=0.12, top=PANEL_TOP, wspace=0.68, hspace=0.82)

    indirect = data["subtype_mediation"][data["subtype_mediation"]["panel"] == "indirect"].copy()
    for i, med in enumerate(MEDIATORS):
        ax = fig.add_subplot(gs[i, 0])
        rows = indirect[indirect["mediator"] == med].copy()
        rows["label"] = rows["subtype"].map({"C1": "C1\nLow-stable", "C2": "C2\nLow-mod/high-fluct", "C3": "C3\nHigh/high-fluct"})
        if i == 0:
            panel_label(ax, "A", -0.14, 1.25)
            ax.set_title("Indirect effects by subtype", loc="left", pad=14, fontweight="bold")
        ax.text(0, 1.03, MEDIATOR_LABELS[med], transform=ax.transAxes, color=PALETTE["ink"], fontsize=8, fontweight="bold")
        draw_interval(ax, rows, "label", "subtype", 0, 0.025, [0, 0.005, 0.010, 0.015, 0.020], "Indirect effect (a x b)")

    ax_b = fig.add_subplot(gs[:, 1])
    contrasts = data["subtype_contrasts"].copy()
    order = []
    for med in MEDIATORS:
        for contrast in ["C2 - C1", "C3 - C1", "C3 - C2"]:
            order.append((med, contrast))
    contrast_rows = []
    for med, contrast in order:
        row = contrasts[(contrasts["mediator"] == med) & (contrasts["contrast"] == contrast)].iloc[0].to_dict()
        row["label"] = f"{contrast} | {MEDIATOR_LABELS[med]}"
        row["color"] = subtype_color(contrast)
        contrast_rows.append(row)
    plot_df = pd.DataFrame(contrast_rows)
    panel_label(ax_b, "B", -0.12, 1.06)
    ax_b.set_title("Subtype contrasts on indirect effects", loc="left", pad=12, fontweight="bold")
    draw_interval(ax_b, plot_df, "label", None, -0.005, 0.025, [-0.005, 0, 0.005, 0.010, 0.015, 0.020], "Contrast in indirect effect", zero=True)
    fig.text(0.64, 0.035, "Only belongingness: credible C3 - C2 contrast.", fontsize=7.2, fontweight="bold", color=PALETTE["ink"])
    save_figure(fig, "Figure_3_subtype_specific_mediation")


def draw_path_decomposition_panel(
    ax,
    df: pd.DataFrame,
    panel: str,
    title: str,
    xmin: float,
    xmax: float,
    xticks: list[float],
    show_labels: bool = False,
) -> None:
    ax.set_title(title, loc="left", pad=8, fontweight="bold")
    ax.set_xlim(xmin, xmax)
    ax.set_xticks(xticks)
    ax.set_xlabel("Posterior median" if panel != "indirect" else "Indirect effect (a x b)")
    if xmin < 0 < xmax:
        ax.axvline(0, color=PALETTE["zero"], lw=0.65, zorder=0)

    y_values = []
    tick_labels = []
    for med_i, med in enumerate(MEDIATORS):
        base = (len(MEDIATORS) - med_i - 1) * 4
        for sub_i, (code, _) in enumerate(SUBTYPES):
            y = base + (len(SUBTYPES) - sub_i - 1)
            row = df[(df["panel"] == panel) & (df["mediator"] == med) & (df["subtype"] == code)].iloc[0]
            ax.errorbar(
                row["estimate"],
                y,
                xerr=[[row["estimate"] - row["low"]], [row["high"] - row["estimate"]]],
                fmt="o",
                color=PALETTE[code],
                ecolor=PALETTE[code],
                elinewidth=1.25,
                capsize=2.6,
                markersize=4.2,
                markeredgecolor="white",
                markeredgewidth=0.5,
                zorder=3,
            )
            y_values.append(y)
            label = f"{MEDIATOR_LABELS[med]}\n{code}" if code == "C1" else code
            tick_labels.append(label)

    ax.set_ylim(-0.7, 10.7)
    ax.set_yticks(y_values)
    ax.set_yticklabels(tick_labels if show_labels else [""] * len(y_values), color=PALETTE["ink"])
    ax.grid(axis="y", color=PALETTE["grid"], lw=0.55)
    clean_axis(ax)


def figure_s1(data: dict[str, pd.DataFrame]) -> None:
    fig = plt.figure(figsize=(WIDTH_IN, HEIGHT_IN), constrained_layout=False)
    gs = fig.add_gridspec(1, 3, width_ratios=[1.14, 1, 1])
    fig.subplots_adjust(left=0.16, right=0.98, bottom=0.12, top=0.86, wspace=0.38)
    axes = [fig.add_subplot(gs[0, i]) for i in range(3)]
    panels = [
        ("a-path", "a-paths: stress to mediator", 0.02, 0.12, [0.02, 0.06, 0.10]),
        ("b-path", "b-paths: mediator to SI", 0.0, 0.30, [0.0, 0.10, 0.20, 0.30]),
        ("indirect", "Indirect effects", 0.0, 0.025, [0.0, 0.01, 0.02]),
    ]
    for i, (ax, (panel, title, xmin, xmax, ticks)) in enumerate(zip(axes, panels)):
        panel_label(ax, chr(ord("A") + i), -0.14, 1.08)
        draw_path_decomposition_panel(ax, data["subtype_mediation"], panel, title, xmin, xmax, ticks, show_labels=i == 0)
    handles = [mpl.lines.Line2D([], [], marker="o", ls="", color=PALETTE[c], label=f"{c}: {l.replace(' / ', '/')}", markersize=5) for c, l in SUBTYPES]
    fig.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.58, 0.975), ncol=3, fontsize=6.6, handletextpad=0.25, columnspacing=1.0)
    save_figure(fig, "Figure_S1_path_decomposition")


def dumbbell(ax, df: pd.DataFrame, labels: list[str], xmin: float, xmax: float, xticks: list[float], show_ylabel: bool = True) -> None:
    ax.set_xlim(xmin, xmax)
    yvals = np.arange(len(labels))[::-1]
    ax.set_yticks(yvals)
    ax.set_yticklabels(labels if show_ylabel else [""] * len(labels), fontweight="bold")
    ax.set_xticks(xticks)
    ax.grid(axis="y", color=PALETTE["grid"], lw=0.55)
    if xmin < 0 < xmax:
        ax.axvline(0, color=PALETTE["zero"], lw=0.65, zorder=0)
    for y, label in zip(yvals, labels):
        before = df[(df["label"] == label) & (df["model"] == "Before mediators")].iloc[0]
        after = df[(df["label"] == label) & (df["model"] == "After mediators")].iloc[0]
        atten = df[(df["label"] == label) & (df["model"] == "Attenuation")].iloc[0]
        color = subtype_color(label)
        ax.plot([before["estimate"], after["estimate"]], [y, y], color=color, alpha=0.35, lw=2.0, zorder=1)
        for row, marker, fill in [(before, "o", "white"), (after, "s", color)]:
            ax.errorbar(
                row["estimate"],
                y,
                xerr=[[row["estimate"] - row["low"]], [row["high"] - row["estimate"]]],
                fmt=marker,
                color=color,
                ecolor=color,
                elinewidth=1.15,
                capsize=2.5,
                markersize=4.0,
                markerfacecolor=fill,
                markeredgecolor=color,
                markeredgewidth=0.9,
                zorder=3,
            )
        ax.text(xmax + (xmax - xmin) * 0.02, y, f"Atten. {atten['estimate']:.4f}", color=PALETTE["muted"], fontsize=6, va="center", fontweight="bold")
    clean_axis(ax)


def figure4(data: dict[str, pd.DataFrame]) -> None:
    fig = plt.figure(figsize=(WIDTH_IN, HEIGHT_IN), constrained_layout=False)
    gs = fig.add_gridspec(2, 2, height_ratios=[1, 1.05])
    fig.subplots_adjust(left=0.08, right=0.91, bottom=0.08, top=0.86, wspace=0.46, hspace=0.62)
    ax_a = fig.add_subplot(gs[0, 0])
    ax_b = fig.add_subplot(gs[0, 1])
    ax_c = fig.add_subplot(gs[1, :])
    legend_handles = [
        mpl.lines.Line2D([], [], marker="o", ls="", color=PALETTE["ink"], markerfacecolor="white", markeredgecolor=PALETTE["ink"], label="Before mediators", markersize=5),
        mpl.lines.Line2D([], [], marker="s", ls="", color=PALETTE["ink"], markerfacecolor=PALETTE["ink"], markeredgecolor=PALETTE["ink"], label="After mediators", markersize=5),
    ]
    fig.legend(handles=legend_handles, loc="upper right", bbox_to_anchor=(0.98, 0.965), ncol=2, fontsize=6.5, handletextpad=0.3)

    panel_label(ax_a, "A", -0.14, 1.12)
    ax_a.set_title("Within-class stress -> SI slopes", loc="left", pad=10, fontweight="bold")
    dumbbell(ax_a, data["attenuation_within"], ["C1", "C2", "C3"], 0, 0.05, [0, 0.01, 0.02, 0.03, 0.04, 0.05])

    panel_label(ax_b, "B", -0.14, 1.12)
    ax_b.set_title("Subtype contrasts", loc="left", pad=10, fontweight="bold")
    dumbbell(ax_b, data["attenuation_contrasts"], ["C2 - C1", "C3 - C1", "C3 - C2"], -0.005, 0.05, [0, 0.01, 0.02, 0.03, 0.04])

    combined = pd.concat(
        [
            data["attenuation_within"][data["attenuation_within"]["model"] == "Attenuation"],
            data["attenuation_contrasts"][data["attenuation_contrasts"]["model"] == "Attenuation"],
        ],
        ignore_index=True,
    )
    combined["color"] = combined["label"].map(lambda value: subtype_color(str(value)))
    panel_label(ax_c, "C", -0.06, 1.12)
    ax_c.set_title("Attenuation after adding mediators", loc="left", pad=10, fontweight="bold")
    draw_interval(ax_c, combined, "label", None, -0.005, 0.04, [-0.005, 0, 0.01, 0.02, 0.03, 0.04], "Before - after mediator adjustment", zero=True)
    ax_c.text(0.55, 1.05, "C2/C3 slopes remain positive but shrink sharply after adjustment.", transform=ax_c.transAxes, fontsize=7, color=PALETTE["muted"])
    save_figure(fig, "Figure_4_stress_si_attenuation")


def validate_outputs(data: dict[str, pd.DataFrame]) -> None:
    checks = {
        "figure1_raw_mean_rows": len(data["profile"]),
        "figure1_prevalence_rows": len(data["prevalence"]),
        "figure2_rows": len(data["overall_mediation"]),
        "figure3_indirect_rows": len(data["subtype_mediation"][data["subtype_mediation"]["panel"] == "indirect"]),
        "figure3_contrast_rows": len(data["subtype_contrasts"]),
        "figure_s1_path_rows": len(data["subtype_mediation"]),
        "figure4_within_rows": len(data["attenuation_within"]),
        "figure4_contrast_rows": len(data["attenuation_contrasts"]),
    }
    expected = {
        "figure1_raw_mean_rows": 15,
        "figure1_prevalence_rows": 6,
        "figure2_rows": 12,
        "figure3_indirect_rows": 9,
        "figure3_contrast_rows": 9,
        "figure_s1_path_rows": 27,
        "figure4_within_rows": 9,
        "figure4_contrast_rows": 9,
    }
    failures = {k: {"expected": expected[k], "observed": v} for k, v in checks.items() if expected[k] != v}
    files = []
    for stem in [
        FIGURE1_STEM,
        "Figure_2_overall_mediation",
        "Figure_3_subtype_specific_mediation",
        "Figure_S1_path_decomposition",
        "Figure_4_stress_si_attenuation",
    ]:
        for ext in [".svg", ".pdf", ".png", ".tiff"]:
            p = OUT_DIR / f"{stem}{ext}"
            files.append({"file": p.name, "exists": p.exists(), "bytes": p.stat().st_size if p.exists() else 0})
            if not p.exists() or p.stat().st_size < 1000:
                failures[p.name] = {"expected": "existing file > 1000 bytes", "observed": p.stat().st_size if p.exists() else "missing"}
        img = Image.open(OUT_DIR / f"{stem}.png")
        extrema = img.convert("L").getextrema()
        files.append({"file": f"{stem}.png", "size": img.size, "gray_extrema": extrema})
        if extrema[0] == extrema[1]:
            failures[f"{stem}.png"] = {"expected": "nonblank image", "observed": extrema}
    for source_name in [
        f"{FIGURE1_STEM}_source_data.csv",
        "Figure_2_overall_mediation_source_data.csv",
        "Figure_3_subtype_specific_mediation_source_data.csv",
        "Figure_S1_path_decomposition_source_data.csv",
        "Figure_4_stress_si_attenuation_source_data.csv",
    ]:
        p = OUT_DIR / source_name
        files.append({"file": p.name, "exists": p.exists(), "bytes": p.stat().st_size if p.exists() else 0})
        if not p.exists() or p.stat().st_size < 100:
            failures[p.name] = {"expected": "existing source data > 100 bytes", "observed": p.stat().st_size if p.exists() else "missing"}
    for ext in [".svg", ".pdf", ".png", ".tiff"]:
        legacy = OUT_DIR / f"{LEGACY_FIGURE1_STEM}{ext}"
        files.append({"file": legacy.name, "exists": legacy.exists(), "legacy": True})
        if legacy.exists():
            failures[legacy.name] = {"expected": "legacy Figure 1 removed", "observed": "still exists"}
    legacy_source = OUT_DIR / f"{LEGACY_FIGURE1_STEM}_source_data.csv"
    files.append({"file": legacy_source.name, "exists": legacy_source.exists(), "legacy": True})
    if legacy_source.exists():
        failures[legacy_source.name] = {"expected": "legacy Figure 1 source data removed", "observed": "still exists"}
    qa = {
        "backend": "matplotlib",
        "matplotlib_version": mpl.__version__,
        "palette": {code: PALETTE[code] for code, _ in SUBTYPES},
        "checks": checks,
        "files": files,
        "failures": failures,
    }
    (OUT_DIR / "generation_QA.json").write_text(json.dumps(qa, indent=2), encoding="utf-8")
    if failures:
        raise RuntimeError(json.dumps(failures, indent=2))


def main() -> None:
    configure_style()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    cleanup_legacy_outputs()
    data = extract_data()
    save_source_data(data)
    figure1(data)
    figure2(data)
    figure3(data)
    figure_s1(data)
    figure4(data)
    validate_outputs(data)
    print(json.dumps({"status": "ok", "backend": "matplotlib", "matplotlib": mpl.__version__, "output_dir": str(OUT_DIR)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
