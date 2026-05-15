"""Dump every cell of Study3_All_Tables.xlsx to a single markdown file."""
import openpyxl, pathlib, sys

src = pathlib.Path(r"D:\心理学\【0428小论文】\Table\Study3_All_Tables.xlsx")
out = pathlib.Path(r"D:\心理学\【0428小论文】\Manuscript_JAD\_xlsx_full_dump.md")
wb = openpyxl.load_workbook(src, data_only=True)

lines = [f"# Full dump of {src.name}\n"]
for s in wb.sheetnames:
    ws = wb[s]
    lines.append(f"\n## {s}  (rows={ws.max_row}, cols={ws.max_column})\n")
    for row in ws.iter_rows(values_only=True):
        cells = [("" if v is None else str(v)) for v in row]
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")
out.write_text("\n".join(lines), encoding="utf-8")
print(f"wrote {out} ({out.stat().st_size} bytes)")
