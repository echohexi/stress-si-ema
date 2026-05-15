from __future__ import annotations

import csv
import json
import math
import re
from dataclasses import dataclass
from html import escape
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.pdfgen import canvas as pdf_canvas


ROOT = Path(__file__).resolve().parents[1]
INPUT_XLSX = ROOT / "Table" / "Study3_All_Tables.xlsx"
OUT_DIR = ROOT / "Figure" / "Study3_main"

W = 1800
H = 1250

COLORS = {
    "ink": "#222222",
    "muted": "#666666",
    "light": "#F3F5F7",
    "grid": "#D7DCE2",
    "zero": "#9EA7B3",
    "C1": "#0072B2",
    "C2": "#E69F00",
    "C3": "#D55E00",
    "Entrapment": "#0072B2",
    "Perceived burdensomeness": "#009E73",
    "Thwarted belongingness": "#CC79A7",
    "before": "#8A8F98",
    "after": "#222222",
    "attenuation": "#D55E00",
}

SUBTYPES = [
    ("C1", "Low-stable"),
    ("C2", "Low-mod / high-fluct"),
    ("C3", "High / high-fluct"),
]

MEDIATORS = [
    "Entrapment",
    "Perceived burdensomeness",
    "Thwarted belongingness",
]


def normalize_minus(value: str) -> str:
    return value.replace("\u2212", "-").replace("\u2013", "-").replace("\u2014", "-")


def parse_float(value) -> float:
    if value is None:
        return math.nan
    return float(str(value).replace(",", "").strip())


EST_RE = re.compile(
    r"^\s*([+-]?\d+(?:\.\d+)?)\s*\[\s*([+-]?\d+(?:\.\d+)?)\s*,\s*([+-]?\d+(?:\.\d+)?)\s*\]\s*(\*)?\s*$"
)


def parse_est_ci(value: str) -> dict[str, float | bool]:
    text = normalize_minus(str(value).strip())
    match = EST_RE.match(text)
    if not match:
        raise ValueError(f"Cannot parse estimate interval: {value!r}")
    est, low, high, star = match.groups()
    return {
        "estimate": float(est),
        "low": float(low),
        "high": float(high),
        "credible": bool(star) or (float(low) > 0) or (float(high) < 0),
    }


def parse_mean_sd(value: str) -> tuple[float, float]:
    text = str(value).replace(",", "").strip()
    match = re.match(r"^([+-]?\d+(?:\.\d+)?)\s*\(([+-]?\d+(?:\.\d+)?)\)$", text)
    if not match:
        raise ValueError(f"Cannot parse mean SD: {value!r}")
    return float(match.group(1)), float(match.group(2))


def parse_n_percent(value: str) -> tuple[int, float | None]:
    text = str(value).replace(",", "").strip()
    match = re.match(r"^(\d+)(?:\s*\(([+-]?\d+(?:\.\d+)?)\))?$", text)
    if not match:
        raise ValueError(f"Cannot parse n percent: {value!r}")
    return int(match.group(1)), float(match.group(2)) if match.group(2) else None


def rows_for_sheet(wb, sheet: str) -> list[list]:
    ws = wb[sheet]
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(v is not None for v in row):
            rows.append(list(row))
    return rows


def safe_name(text: str) -> str:
    return (
        text.replace(" ", "_")
        .replace("/", "_")
        .replace("(", "")
        .replace(")", "")
        .replace(">", "gt")
        .replace("<", "lt")
        .replace("'", "")
    )


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        return
    fields = []
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def hex_to_rgb(color: str) -> tuple[int, int, int]:
    color = color.lstrip("#")
    return tuple(int(color[i : i + 2], 16) for i in (0, 2, 4))


def blend(c1: str, c2: str, t: float) -> str:
    r1, g1, b1 = hex_to_rgb(c1)
    r2, g2, b2 = hex_to_rgb(c2)
    return f"#{round(r1 + (r2 - r1) * t):02x}{round(g1 + (g2 - g1) * t):02x}{round(b1 + (b2 - b1) * t):02x}"


def scale_value(value: float, min_value: float, max_value: float, a: float, b: float) -> float:
    if max_value == min_value:
        return (a + b) / 2
    return a + (value - min_value) / (max_value - min_value) * (b - a)


@dataclass
class Op:
    kind: str
    args: tuple
    kwargs: dict


class Figure:
    def __init__(self, width: int = W, height: int = H):
        self.width = width
        self.height = height
        self.ops: list[Op] = []

    def line(self, x1, y1, x2, y2, color=COLORS["ink"], width=2, dash=None):
        self.ops.append(Op("line", (x1, y1, x2, y2), {"color": color, "width": width, "dash": dash}))

    def rect(self, x, y, w, h, fill="none", stroke=COLORS["ink"], width=1, radius=0):
        self.ops.append(Op("rect", (x, y, w, h), {"fill": fill, "stroke": stroke, "width": width, "radius": radius}))

    def circle(self, x, y, r, fill=COLORS["ink"], stroke="none", width=1):
        self.ops.append(Op("circle", (x, y, r), {"fill": fill, "stroke": stroke, "width": width}))

    def polygon(self, points, fill=COLORS["ink"], stroke="none", width=1):
        self.ops.append(Op("polygon", (points,), {"fill": fill, "stroke": stroke, "width": width}))

    def text(self, x, y, text, size=24, color=COLORS["ink"], anchor="start", bold=False):
        self.ops.append(Op("text", (x, y, str(text)), {"size": size, "color": color, "anchor": anchor, "bold": bold}))

    def save(self, base: Path):
        base.parent.mkdir(parents=True, exist_ok=True)
        self._save_svg(base.with_suffix(".svg"))
        self._save_png(base.with_suffix(".png"), scale=1, dpi=300)
        self._save_png(base.with_suffix(".tiff"), scale=2, dpi=600)
        self._save_pdf(base.with_suffix(".pdf"))

    def _save_svg(self, path: Path):
        out = [
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{self.width}" height="{self.height}" viewBox="0 0 {self.width} {self.height}">',
            '<rect width="100%" height="100%" fill="white"/>',
            '<style>text{font-family:Arial,Helvetica,sans-serif;dominant-baseline:middle;}</style>',
        ]
        for op in self.ops:
            kw = op.kwargs
            if op.kind == "line":
                x1, y1, x2, y2 = op.args
                dash = f' stroke-dasharray="{kw["dash"]}"' if kw.get("dash") else ""
                out.append(
                    f'<line x1="{x1:.2f}" y1="{y1:.2f}" x2="{x2:.2f}" y2="{y2:.2f}" stroke="{kw["color"]}" stroke-width="{kw["width"]}"{dash} stroke-linecap="round"/>'
                )
            elif op.kind == "rect":
                x, y, w, h = op.args
                out.append(
                    f'<rect x="{x:.2f}" y="{y:.2f}" width="{w:.2f}" height="{h:.2f}" rx="{kw["radius"]}" fill="{kw["fill"]}" stroke="{kw["stroke"]}" stroke-width="{kw["width"]}"/>'
                )
            elif op.kind == "circle":
                x, y, r = op.args
                out.append(
                    f'<circle cx="{x:.2f}" cy="{y:.2f}" r="{r:.2f}" fill="{kw["fill"]}" stroke="{kw["stroke"]}" stroke-width="{kw["width"]}"/>'
                )
            elif op.kind == "polygon":
                (points,) = op.args
                pts = " ".join(f"{x:.2f},{y:.2f}" for x, y in points)
                out.append(f'<polygon points="{pts}" fill="{kw["fill"]}" stroke="{kw["stroke"]}" stroke-width="{kw["width"]}"/>')
            elif op.kind == "text":
                x, y, text = op.args
                weight = "700" if kw.get("bold") else "400"
                anchor = {"start": "start", "middle": "middle", "end": "end"}.get(kw["anchor"], "start")
                out.append(
                    f'<text x="{x:.2f}" y="{y:.2f}" font-size="{kw["size"]}" fill="{kw["color"]}" text-anchor="{anchor}" font-weight="{weight}">{escape(text)}</text>'
                )
        out.append("</svg>")
        path.write_text("\n".join(out), encoding="utf-8")

    def _font(self, size: int, bold: bool = False, scale: int = 1):
        candidates = [
            Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
            Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
        ]
        for cand in candidates:
            if cand.exists():
                return ImageFont.truetype(str(cand), int(size * scale))
        return ImageFont.load_default()

    def _save_png(self, path: Path, scale: int, dpi: int):
        img = Image.new("RGB", (self.width * scale, self.height * scale), "white")
        draw = ImageDraw.Draw(img)

        def s(v):
            return int(round(v * scale))

        for op in self.ops:
            kw = op.kwargs
            if op.kind == "line":
                x1, y1, x2, y2 = op.args
                draw.line((s(x1), s(y1), s(x2), s(y2)), fill=kw["color"], width=max(1, s(kw["width"])))
            elif op.kind == "rect":
                x, y, w, h = op.args
                box = (s(x), s(y), s(x + w), s(y + h))
                fill = None if kw["fill"] == "none" else kw["fill"]
                outline = None if kw["stroke"] == "none" else kw["stroke"]
                if kw.get("radius", 0):
                    draw.rounded_rectangle(box, radius=s(kw["radius"]), fill=fill, outline=outline, width=max(1, s(kw["width"])))
                else:
                    draw.rectangle(box, fill=fill, outline=outline, width=max(1, s(kw["width"])))
            elif op.kind == "circle":
                x, y, r = op.args
                box = (s(x - r), s(y - r), s(x + r), s(y + r))
                fill = None if kw["fill"] == "none" else kw["fill"]
                outline = None if kw["stroke"] == "none" else kw["stroke"]
                draw.ellipse(box, fill=fill, outline=outline, width=max(1, s(kw["width"])))
            elif op.kind == "polygon":
                (points,) = op.args
                fill = None if kw["fill"] == "none" else kw["fill"]
                outline = None if kw["stroke"] == "none" else kw["stroke"]
                draw.polygon([(s(x), s(y)) for x, y in points], fill=fill, outline=outline)
            elif op.kind == "text":
                x, y, text = op.args
                font = self._font(kw["size"], kw.get("bold", False), scale)
                bbox = draw.textbbox((0, 0), text, font=font)
                tw = bbox[2] - bbox[0]
                th = bbox[3] - bbox[1]
                tx = s(x)
                if kw["anchor"] == "middle":
                    tx -= tw // 2
                elif kw["anchor"] == "end":
                    tx -= tw
                ty = s(y) - th // 2
                draw.text((tx, ty), text, fill=kw["color"], font=font)
        img.save(path, dpi=(dpi, dpi))

    def _save_pdf(self, path: Path):
        width_pt = self.width / 250 * 72
        height_pt = self.height / 250 * 72
        sx = width_pt / self.width
        sy = height_pt / self.height
        c = pdf_canvas.Canvas(str(path), pagesize=(width_pt, height_pt))

        def pdf_color(value: str):
            if value == "white":
                return colors.white
            if value == "black":
                return colors.black
            return colors.HexColor(value)

        def conv(x, y):
            return x * sx, height_pt - y * sy

        def set_stroke(hex_color, width):
            if hex_color != "none":
                c.setStrokeColor(pdf_color(hex_color))
            c.setLineWidth(width * sx)

        for op in self.ops:
            kw = op.kwargs
            if op.kind == "line":
                x1, y1, x2, y2 = op.args
                set_stroke(kw["color"], kw["width"])
                if kw.get("dash"):
                    c.setDash([4 * sx, 4 * sx])
                else:
                    c.setDash([])
                c.line(*conv(x1, y1), *conv(x2, y2))
            elif op.kind == "rect":
                x, y, w, h = op.args
                c.setFillColor(colors.white if kw["fill"] == "none" else pdf_color(kw["fill"]))
                c.setStrokeColor(colors.white if kw["stroke"] == "none" else pdf_color(kw["stroke"]))
                c.setLineWidth(kw["width"] * sx)
                px, py = conv(x, y + h)
                c.rect(px, py, w * sx, h * sy, fill=kw["fill"] != "none", stroke=kw["stroke"] != "none")
            elif op.kind == "circle":
                x, y, r = op.args
                c.setFillColor(colors.white if kw["fill"] == "none" else pdf_color(kw["fill"]))
                c.setStrokeColor(colors.white if kw["stroke"] == "none" else pdf_color(kw["stroke"]))
                c.setLineWidth(kw["width"] * sx)
                px, py = conv(x, y)
                c.circle(px, py, r * sx, fill=kw["fill"] != "none", stroke=kw["stroke"] != "none")
            elif op.kind == "polygon":
                (points,) = op.args
                path_obj = c.beginPath()
                x0, y0 = conv(*points[0])
                path_obj.moveTo(x0, y0)
                for point in points[1:]:
                    path_obj.lineTo(*conv(*point))
                path_obj.close()
                c.setFillColor(colors.white if kw["fill"] == "none" else pdf_color(kw["fill"]))
                c.setStrokeColor(colors.white if kw["stroke"] == "none" else pdf_color(kw["stroke"]))
                c.drawPath(path_obj, fill=kw["fill"] != "none", stroke=kw["stroke"] != "none")
            elif op.kind == "text":
                x, y, text = op.args
                font = "Helvetica-Bold" if kw.get("bold") else "Helvetica"
                size = kw["size"] * sx
                c.setFont(font, size)
                c.setFillColor(pdf_color(kw["color"]))
                px, py = conv(x, y)
                sw = c.stringWidth(text, font, size)
                if kw["anchor"] == "middle":
                    px -= sw / 2
                elif kw["anchor"] == "end":
                    px -= sw
                c.drawString(px, py - size * 0.33, text)
        c.save()


def panel_label(fig: Figure, label: str, x: float, y: float):
    fig.text(x, y, label, size=38, bold=True)


def title(fig: Figure, text: str, x: float, y: float):
    fig.text(x, y, text, size=30, bold=True)


def subtitle(fig: Figure, text: str, x: float, y: float):
    fig.text(x, y, text, size=20, color=COLORS["muted"])


def wrapped_text(fig: Figure, x: float, y: float, text: str, max_chars: int, size=18, color=COLORS["ink"], bold=False, line_height=26):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if len(trial) > max_chars and current:
            lines.append(current)
            current = word
        else:
            current = trial
    if current:
        lines.append(current)
    for i, line in enumerate(lines):
        fig.text(x, y + i * line_height, line, size=size, color=color, bold=bold)


def arrow(fig: Figure, x1, y1, x2, y2, color=COLORS["ink"], width=4):
    fig.line(x1, y1, x2, y2, color=color, width=width)
    ang = math.atan2(y2 - y1, x2 - x1)
    size = 18
    p1 = (x2, y2)
    p2 = (x2 - size * math.cos(ang - 0.45), y2 - size * math.sin(ang - 0.45))
    p3 = (x2 - size * math.cos(ang + 0.45), y2 - size * math.sin(ang + 0.45))
    fig.polygon([p1, p2, p3], fill=color)


def draw_x_axis(fig: Figure, x0, y0, w, xmin, xmax, ticks, label=None, zero=False):
    if zero and xmin < 0 < xmax:
        zx = scale_value(0, xmin, xmax, x0, x0 + w)
        fig.line(zx, y0 - 18, zx, y0 - 330, color=COLORS["zero"], width=2, dash="7 6")
    fig.line(x0, y0, x0 + w, y0, color=COLORS["ink"], width=2)
    for t in ticks:
        x = scale_value(t, xmin, xmax, x0, x0 + w)
        fig.line(x, y0, x, y0 + 10, color=COLORS["ink"], width=2)
        fig.text(x, y0 + 34, f"{t:g}", size=17, color=COLORS["muted"], anchor="middle")
    if label:
        fig.text(x0 + w / 2, y0 + 75, label, size=19, color=COLORS["muted"], anchor="middle")


def interval_plot(
    fig: Figure,
    x,
    y,
    w,
    h,
    rows,
    xmin,
    xmax,
    ticks,
    xlabel,
    color_key=None,
    zero=False,
    label_width=315,
    marker_size=10,
):
    plot_x = x + label_width
    plot_w = w - label_width - 20
    row_gap = h / max(1, len(rows))
    axis_y = y + h + 18
    if zero and xmin < 0 < xmax:
        zx = scale_value(0, xmin, xmax, plot_x, plot_x + plot_w)
        fig.line(zx, y - 12, zx, axis_y, color=COLORS["zero"], width=2, dash="7 6")
    for idx, row in enumerate(rows):
        yy = y + row_gap * idx + row_gap / 2
        fig.line(plot_x, yy, plot_x + plot_w, yy, color=COLORS["light"], width=2)
        fig.text(plot_x - 18, yy, row["label"], size=18, color=COLORS["ink"], anchor="end")
        color = COLORS.get(row.get(color_key, ""), COLORS["ink"]) if color_key else row.get("color", COLORS["ink"])
        x_low = scale_value(row["low"], xmin, xmax, plot_x, plot_x + plot_w)
        x_high = scale_value(row["high"], xmin, xmax, plot_x, plot_x + plot_w)
        x_est = scale_value(row["estimate"], xmin, xmax, plot_x, plot_x + plot_w)
        fig.line(x_low, yy, x_high, yy, color=color, width=4)
        fig.line(x_low, yy - 10, x_low, yy + 10, color=color, width=3)
        fig.line(x_high, yy - 10, x_high, yy + 10, color=color, width=3)
        fig.circle(x_est, yy, marker_size, fill=color, stroke="white", width=3)
    draw_x_axis(fig, plot_x, axis_y, plot_w, xmin, xmax, ticks, xlabel, zero=False)


def extract_data():
    wb = load_workbook(INPUT_XLSX, data_only=True)
    data = {}

    t1 = rows_for_sheet(wb, "Table 1")
    by_name = {r[0]: r for r in t1 if r and r[0]}
    sample_rows = [
        ("Participants", "Participants with >=1 person-wave", parse_float(by_name["Unique participants contributing \u22651 person-wave to subtype, n"][1])),
        ("Person-waves", "Person-waves", parse_float(str(by_name["Person-waves, n (% of total)"][1]).replace(",", ""))),
        ("EMA prompts", "EMA prompts", parse_float(str(by_name["EMA prompts, n"][1]).replace(",", ""))),
        ("EMA waves", "EMA waves", 4),
    ]
    data["sample"] = [
        {"panel": "sample", "metric": metric, "label": label, "value": value}
        for metric, label, value in sample_rows
    ]
    comp = []
    total_waves = parse_float(str(by_name["Person-waves, n (% of total)"][1]).replace(",", ""))
    total_prompts = parse_float(str(by_name["EMA prompts, n"][1]).replace(",", ""))
    for i, (code, label) in enumerate(SUBTYPES, start=2):
        n, pct = parse_n_percent(by_name["Person-waves, n (% of total)"][i])
        comp.append({"panel": "composition", "measure": "Person-waves", "subtype": code, "subtype_label": label, "n": n, "percent": pct})
        n_prompt = parse_float(str(by_name["EMA prompts, n"][i]).replace(",", ""))
        comp.append({"panel": "composition", "measure": "EMA prompts", "subtype": code, "subtype_label": label, "n": int(n_prompt), "percent": n_prompt / total_prompts * 100})
    data["composition"] = comp

    profile_vars = [
        "Stress sum intensity",
        "Entrapment",
        "Perceived burdensomeness",
        "Thwarted belongingness",
        "Suicidal ideation",
    ]
    profile = []
    for var in profile_vars:
        row = by_name[var]
        for i, (code, label) in enumerate(SUBTYPES, start=2):
            mean, sd = parse_mean_sd(row[i])
            profile.append({"panel": "profile", "variable": var, "subtype": code, "subtype_label": label, "mean": mean, "sd": sd})
    data["profile"] = profile

    prevalence_vars = ["Any momentary stressor endorsed", "Any momentary suicidal ideation (SI > 1)"]
    prevalence = []
    for var in prevalence_vars:
        row = by_name[var]
        for i, (code, label) in enumerate(SUBTYPES, start=2):
            prevalence.append({"panel": "prevalence", "variable": var, "subtype": code, "subtype_label": label, "percent": parse_float(row[i])})
    data["prevalence"] = prevalence

    t2 = rows_for_sheet(wb, "Table 2")
    section = None
    overall = []
    for row in t2:
        first = row[0]
        if first in ["a-paths: Stress \u2192 Mediator", "b-paths: Mediator \u2192 Suicidal ideation", "Within-person indirect effects (a \u00d7 b)", "Dominance contrasts among indirect effects"]:
            section = first
            continue
        if section and row[1] is not None and "[" in str(row[2]):
            interval = parse_est_ci(f"{row[1]} {row[2]}")
            if section.startswith("a-paths"):
                group = "a-path"
            elif section.startswith("b-paths"):
                group = "b-path"
            elif section.startswith("Within"):
                group = "indirect"
            else:
                group = "dominance"
            mediator = None
            for med in MEDIATORS:
                if med in first:
                    mediator = med
            overall.append({"panel": group, "path": first, "mediator": mediator or first, "pd": row[3], **interval})
    data["overall_mediation"] = overall

    t3 = rows_for_sheet(wb, "Table 3")
    subtype_rows = []
    contrasts = []
    current_mediator = None
    for row in t3:
        first = row[0]
        if first in MEDIATORS:
            current_mediator = first
            continue
        if not current_mediator or first is None:
            continue
        if str(first).startswith("C"):
            subtype = str(first).split()[0]
            a = parse_est_ci(row[1])
            b = parse_est_ci(row[2])
            ind = parse_est_ci(row[3])
            subtype_rows.append({"panel": "subtype_indirect", "mediator": current_mediator, "subtype": subtype, **ind})
            subtype_rows.append({"panel": "subtype_a_path", "mediator": current_mediator, "subtype": subtype, **a})
            subtype_rows.append({"panel": "subtype_b_path", "mediator": current_mediator, "subtype": subtype, **b})
        elif str(first).startswith("Indirect contrast"):
            contrast = normalize_minus(str(first).replace("Indirect contrast: ", ""))
            ind = parse_est_ci(str(row[3]).replace("*", " *"))
            contrasts.append({"panel": "subtype_contrast", "mediator": current_mediator, "contrast": contrast, **ind})
    data["subtype_mediation"] = subtype_rows
    data["subtype_contrasts"] = contrasts

    t4 = rows_for_sheet(wb, "Table 4")
    slopes = []
    contrasts4 = []
    current_section = None
    for row in t4:
        first = row[0]
        if first == "Within-class stress \u2192 SI slope":
            current_section = "within"
            continue
        if first == "Subtype contrasts on stress \u2192 SI slope":
            current_section = "contrast"
            continue
        if current_section and first and row[1] is not None and "[" in str(row[1]):
            before = parse_est_ci(row[1])
            after = parse_est_ci(row[2])
            atten = parse_est_ci(row[3])
            dest = slopes if current_section == "within" else contrasts4
            label = str(first).split()[0] if current_section == "within" else normalize_minus(str(first))
            for model, vals in [("Before mediators", before), ("After mediators", after), ("Attenuation", atten)]:
                dest.append({"panel": current_section, "label": label, "model": model, **vals})
    data["attenuation_within"] = slopes
    data["attenuation_contrasts"] = contrasts4

    return data


def save_source_data(data):
    fig1 = data["sample"] + data["composition"] + data["profile"] + data["prevalence"]
    write_csv(OUT_DIR / "Figure_1_subtype_landscape_source_data.csv", fig1)
    write_csv(OUT_DIR / "Figure_2_overall_mediation_source_data.csv", data["overall_mediation"])
    write_csv(OUT_DIR / "Figure_3_subtype_specific_mediation_source_data.csv", data["subtype_mediation"] + data["subtype_contrasts"])
    write_csv(OUT_DIR / "Figure_4_stress_si_attenuation_source_data.csv", data["attenuation_within"] + data["attenuation_contrasts"])


def figure1(data):
    fig = Figure(W, H)
    fig.text(70, 46, "Figure 1. Subtype landscape across EMA sampling and momentary profiles", size=34, bold=True)
    subtitle(fig, "Values are descriptive summaries from the manuscript table; heatmap color is scaled within each row.", 70, 84)

    panel_label(fig, "A", 70, 145)
    title(fig, "Study sampling frame", 115, 150)
    x0, y0 = 100, 210
    card_w, card_h = 260, 126
    metrics = data["sample"]
    for i, row in enumerate(metrics):
        x = x0 + (i % 2) * 295
        y = y0 + (i // 2) * 150
        fig.rect(x, y, card_w, card_h, fill=COLORS["light"], stroke="#E1E5EA", width=2, radius=10)
        value = int(row["value"]) if row["value"] == int(row["value"]) else row["value"]
        fig.text(x + 24, y + 36, f"{value:,}", size=35, bold=True, color=COLORS["ink"])
        wrapped_text(fig, x + 24, y + 76, row["label"], 28, size=16, color=COLORS["muted"], line_height=20)

    panel_label(fig, "B", 710, 145)
    title(fig, "Subtype composition", 755, 150)
    subtitle(fig, "Share of person-waves and EMA prompts", 755, 186)
    comp = data["composition"]
    for j, measure in enumerate(["Person-waves", "EMA prompts"]):
        y = 250 + j * 125
        fig.text(760, y + 21, measure, size=21, bold=True)
        bx, bw = 1010, 600
        cur = bx
        for code, label in SUBTYPES:
            row = next(r for r in comp if r["measure"] == measure and r["subtype"] == code)
            seg_w = bw * row["percent"] / 100
            fig.rect(cur, y, seg_w, 42, fill=COLORS[code], stroke="white", width=3)
            if seg_w > 60:
                fig.text(cur + seg_w / 2, y + 21, f"{row['percent']:.1f}%", size=17, color="white", anchor="middle", bold=True)
            cur += seg_w
        fig.rect(bx, y, bw, 42, fill="none", stroke=COLORS["ink"], width=2)
    lx = 1020
    for i, (code, label) in enumerate(SUBTYPES):
        fig.circle(lx + i * 230, 455, 11, fill=COLORS[code])
        fig.text(lx + 18 + i * 230, 455, f"{code}: {label}", size=17, color=COLORS["muted"])

    panel_label(fig, "C", 70, 560)
    title(fig, "Subtype momentary profile", 115, 565)
    profile = data["profile"]
    vars_order = ["Stress sum intensity", "Entrapment", "Perceived burdensomeness", "Thwarted belongingness", "Suicidal ideation"]
    x, y = 115, 630
    cell_w, cell_h = 210, 76
    fig.text(x + 390, y - 34, "C1", size=20, bold=True, anchor="middle")
    fig.text(x + 600, y - 34, "C2", size=20, bold=True, anchor="middle")
    fig.text(x + 810, y - 34, "C3", size=20, bold=True, anchor="middle")
    for r_i, var in enumerate(vars_order):
        row_vals = [next(r for r in profile if r["variable"] == var and r["subtype"] == code)["mean"] for code, _ in SUBTYPES]
        lo, hi = min(row_vals), max(row_vals)
        yy = y + r_i * cell_h
        fig.text(x, yy + cell_h / 2, var.replace("Perceived ", "Perceived\n"), size=18, anchor="start")
        for c_i, (code, _) in enumerate(SUBTYPES):
            value = row_vals[c_i]
            t = 0.18 + 0.78 * ((value - lo) / (hi - lo) if hi > lo else 0.5)
            fill = blend("#FFFFFF", COLORS[code], t)
            xx = x + 285 + c_i * cell_w
            fig.rect(xx, yy, cell_w - 8, cell_h - 8, fill=fill, stroke="white", width=3)
            fig.text(xx + (cell_w - 8) / 2, yy + (cell_h - 8) / 2, f"{value:.2f}", size=22, bold=True, anchor="middle")
    subtitle(fig, "Raw means shown; color intensity is scaled within each variable.", 115, 1060)

    panel_label(fig, "D", 1065, 560)
    title(fig, "Prompt-level prevalence", 1110, 565)
    prev = data["prevalence"]
    plot_x, plot_y, plot_w = 1190, 680, 460
    outcomes = [
        ("Any momentary stressor endorsed", "Any stressor"),
        ("Any momentary suicidal ideation (SI > 1)", "Any SI"),
    ]
    for i, (key, label) in enumerate(outcomes):
        yy = plot_y + i * 150
        fig.text(plot_x - 30, yy, label, size=21, anchor="end", bold=True)
        fig.line(plot_x, yy, plot_x + plot_w, yy, color=COLORS["light"], width=4)
        for code, _ in SUBTYPES:
            row = next(r for r in prev if r["variable"] == key and r["subtype"] == code)
            xx = scale_value(row["percent"], 0, 100, plot_x, plot_x + plot_w)
            fig.circle(xx, yy, 15, fill=COLORS[code], stroke="white", width=3)
            fig.text(xx, yy - 34, f"{row['percent']:.1f}", size=15, anchor="middle", color=COLORS[code], bold=True)
    draw_x_axis(fig, plot_x, 1010, plot_w, 0, 100, [0, 25, 50, 75, 100], "Prompt-level prevalence (%)")
    fig.save(OUT_DIR / "Figure_1_subtype_landscape")


def figure2(data):
    fig = Figure(W, H)
    fig.text(70, 46, "Figure 2. Overall within-person parallel mediation", size=34, bold=True)
    subtitle(fig, "Dots show posterior medians; intervals are 95% Bayesian credible intervals.", 70, 84)
    rows = data["overall_mediation"]

    panel_label(fig, "A", 70, 145)
    title(fig, "Evidence logic", 115, 150)
    sx, sy = 160, 370
    ix, iy = 800, 370
    fig.rect(sx - 70, sy - 50, 190, 95, fill=COLORS["light"], stroke="#E1E5EA", width=2)
    fig.text(sx + 25, sy, "Stress", size=26, bold=True, anchor="middle")
    fig.rect(ix - 50, iy - 50, 190, 95, fill=COLORS["light"], stroke="#E1E5EA", width=2)
    fig.text(ix + 45, iy, "SI", size=26, bold=True, anchor="middle")
    med_y = [245, 370, 495]
    for med, yy in zip(MEDIATORS, med_y):
        mx = 465
        fig.rect(mx - 120, yy - 38, 260, 76, fill="#FFFFFF", stroke=COLORS[med], width=4)
        label = med.replace("Perceived burdensomeness", "Burdensomeness").replace("Thwarted belongingness", "Belongingness")
        fig.text(mx + 10, yy, label, size=19, bold=True, anchor="middle", color=COLORS[med])
        arrow(fig, sx + 120, sy, mx - 122, yy, color=COLORS[med], width=4)
        arrow(fig, mx + 142, yy, ix - 54, sy, color=COLORS[med], width=4)
    fig.text(440, 585, "All three indirect pathways were credibly positive.", size=22, bold=True, anchor="middle")

    a_rows = [{"label": r["mediator"].replace("Perceived ", "Perceived\n").replace("Thwarted ", "Thwarted\n"), **r} for r in rows if r["panel"] == "a-path"]
    b_rows = [{"label": r["mediator"].replace("Perceived ", "Perceived\n").replace("Thwarted ", "Thwarted\n"), **r} for r in rows if r["panel"] == "b-path"]
    ind_rows = [{"label": r["mediator"].replace("Perceived ", "Perceived\n").replace("Thwarted ", "Thwarted\n"), **r} for r in rows if r["panel"] == "indirect"]
    dom_rows = [{"label": r["path"].replace("Perceived burdensomeness", "Burdensomeness").replace("Thwarted belongingness", "Belongingness"), **r} for r in rows if r["panel"] == "dominance"]

    panel_label(fig, "B", 985, 145)
    title(fig, "a-paths: stress to mediator", 1030, 150)
    interval_plot(fig, 1010, 220, 690, 210, a_rows, 0.03, 0.08, [0.03, 0.05, 0.07], None, "mediator", label_width=245)

    panel_label(fig, "C", 985, 530)
    title(fig, "b-paths: mediator to SI", 1030, 535)
    interval_plot(fig, 1010, 605, 690, 210, b_rows, 0.00, 0.065, [0, 0.02, 0.04, 0.06], None, "mediator", label_width=245)

    panel_label(fig, "D", 70, 705)
    title(fig, "Indirect effects", 115, 710)
    interval_plot(fig, 115, 785, 760, 240, ind_rows, 0, 0.0035, [0, 0.001, 0.002, 0.003], "Indirect effect (a x b)", "mediator", label_width=285)

    panel_label(fig, "E", 985, 865)
    title(fig, "Dominance contrasts", 1030, 870)
    interval_plot(fig, 1030, 925, 675, 190, dom_rows, -0.0006, 0.0026, [-0.0005, 0, 0.001, 0.002], "Difference in indirect effect", None, zero=True, label_width=330)
    fig.save(OUT_DIR / "Figure_2_overall_mediation")


def figure3(data):
    fig = Figure(W, H)
    fig.text(70, 46, "Figure 3. Subtype-specific mediation is amplified in fluctuating subtypes", size=33, bold=True)
    subtitle(fig, "Dots show posterior medians; intervals are 95% Bayesian credible intervals.", 70, 84)

    indirect = [r for r in data["subtype_mediation"] if r["panel"] == "subtype_indirect"]
    contrasts = data["subtype_contrasts"]

    panel_label(fig, "A", 70, 145)
    title(fig, "Indirect effects by subtype", 115, 150)
    y_base = 225
    for i, med in enumerate(MEDIATORS):
        y = y_base + i * 285
        fig.text(120, y - 35, med, size=24, bold=True, color=COLORS[med])
        rows = []
        for code, label in SUBTYPES:
            r = next(v for v in indirect if v["mediator"] == med and v["subtype"] == code)
            rows.append({"label": f"{code}: {label}", "subtype": code, **r})
        interval_plot(fig, 130, y, 750, 185, rows, 0, 0.025, [0, 0.005, 0.010, 0.015, 0.020], "Indirect effect (a x b)", "subtype", label_width=260, marker_size=11)

    panel_label(fig, "B", 1000, 145)
    title(fig, "Subtype contrasts on indirect effects", 1045, 150)
    contrast_labels = ["C2 - C1", "C3 - C1", "C3 - C2"]
    rows = []
    for med in MEDIATORS:
        for contrast in contrast_labels:
            r = next(v for v in contrasts if v["mediator"] == med and v["contrast"] == contrast)
            rows.append({
                "label": f"{contrast} | {med.replace('Perceived burdensomeness', 'Burdensomeness').replace('Thwarted belongingness', 'Belongingness')}",
                "color": COLORS[med],
                **r,
            })
    interval_plot(fig, 1035, 235, 710, 695, rows, -0.005, 0.025, [-0.005, 0, 0.005, 0.010, 0.015, 0.020], "Contrast in indirect effect", None, zero=True, label_width=375, marker_size=9)
    fig.rect(1070, 1015, 620, 86, fill=COLORS["light"], stroke="#E1E5EA", width=2)
    wrapped_text(
        fig,
        1095,
        1045,
        "C3 - C2 credibly exceeded zero only for thwarted belongingness, indicating a subtype-specific interpersonal pathway.",
        70,
        size=20,
        color=COLORS["ink"],
        bold=True,
        line_height=25,
    )
    fig.save(OUT_DIR / "Figure_3_subtype_specific_mediation")


def figure4(data):
    fig = Figure(W, H)
    fig.text(70, 46, "Figure 4. Mediators attenuate the within-person stress-to-SI slope", size=33, bold=True)
    subtitle(fig, "Dots show posterior medians; intervals are 95% Bayesian credible intervals.", 70, 84)
    fig.circle(1185, 85, 11, fill=COLORS["before"])
    fig.text(1205, 85, "Before mediators", size=18, color=COLORS["muted"])
    fig.circle(1395, 85, 11, fill=COLORS["after"])
    fig.text(1415, 85, "After mediators", size=18, color=COLORS["muted"])

    within = data["attenuation_within"]
    contrasts = data["attenuation_contrasts"]

    panel_label(fig, "A", 70, 145)
    title(fig, "Within-class stress -> SI slopes", 115, 150)
    dumbbell(fig, 130, 245, 760, 315, within, ["C1", "C2", "C3"], 0, 0.05, [0, 0.01, 0.02, 0.03, 0.04, 0.05])

    panel_label(fig, "B", 990, 145)
    title(fig, "Subtype contrasts", 1035, 150)
    dumbbell(fig, 1035, 245, 680, 315, contrasts, ["C2 - C1", "C3 - C1", "C3 - C2"], -0.005, 0.05, [0, 0.01, 0.02, 0.03, 0.04], zero=True)

    panel_label(fig, "C", 70, 660)
    title(fig, "Attenuation after adding mediators", 115, 665)
    att_rows = []
    for label in ["C1", "C2", "C3"]:
        r = next(v for v in within if v["label"] == label and v["model"] == "Attenuation")
        att_rows.append({"label": label, "color": COLORS.get(label, COLORS["attenuation"]), **r})
    for label in ["C2 - C1", "C3 - C1", "C3 - C2"]:
        r = next(v for v in contrasts if v["label"] == label and v["model"] == "Attenuation")
        att_rows.append({"label": label, "color": COLORS["attenuation"], **r})
    interval_plot(fig, 140, 750, 1460, 310, att_rows, -0.005, 0.04, [-0.005, 0, 0.01, 0.02, 0.03, 0.04], "Before - after mediator adjustment", None, zero=True, label_width=190, marker_size=10)
    fig.text(1295, 705, "C2/C3 slopes remain positive but shrink sharply after adjustment.", size=20, color=COLORS["muted"], anchor="middle")
    fig.save(OUT_DIR / "Figure_4_stress_si_attenuation")


def dumbbell(fig, x, y, w, h, data_rows, labels, xmin, xmax, ticks, zero=False):
    label_w = 175
    plot_x = x + label_w
    plot_w = w - label_w - 25
    gap = h / len(labels)
    axis_y = y + h + 25
    if zero and xmin < 0 < xmax:
        zx = scale_value(0, xmin, xmax, plot_x, plot_x + plot_w)
        fig.line(zx, y - 15, zx, axis_y, color=COLORS["zero"], width=2, dash="7 6")
    for i, label in enumerate(labels):
        yy = y + i * gap + gap / 2
        fig.text(plot_x - 20, yy, label, size=21, bold=True, anchor="end")
        before = next(v for v in data_rows if v["label"] == label and v["model"] == "Before mediators")
        after = next(v for v in data_rows if v["label"] == label and v["model"] == "After mediators")
        bx = scale_value(before["estimate"], xmin, xmax, plot_x, plot_x + plot_w)
        ax = scale_value(after["estimate"], xmin, xmax, plot_x, plot_x + plot_w)
        fig.line(plot_x, yy, plot_x + plot_w, yy, color=COLORS["light"], width=4)
        fig.line(bx, yy, ax, yy, color="#B6BEC9", width=6)
        for row, color, offset in [(before, COLORS["before"], -15), (after, COLORS["after"], 15)]:
            lo = scale_value(row["low"], xmin, xmax, plot_x, plot_x + plot_w)
            hi = scale_value(row["high"], xmin, xmax, plot_x, plot_x + plot_w)
            est = scale_value(row["estimate"], xmin, xmax, plot_x, plot_x + plot_w)
            fig.line(lo, yy + offset, hi, yy + offset, color=color, width=3)
            fig.circle(est, yy + offset, 10, fill=color, stroke="white", width=3)
        att = next(v for v in data_rows if v["label"] == label and v["model"] == "Attenuation")
        fig.text(plot_x + plot_w + 10, yy, f"Atten. {att['estimate']:.4f}", size=16, color=COLORS["attenuation"], bold=True)
    draw_x_axis(fig, plot_x, axis_y, plot_w, xmin, xmax, ticks, None, zero=False)


def validate_outputs(data):
    checks = {
        "figure1_profile_rows": len(data["profile"]),
        "figure1_prevalence_rows": len(data["prevalence"]),
        "figure2_rows": len(data["overall_mediation"]),
        "figure3_indirect_rows": len([r for r in data["subtype_mediation"] if r["panel"] == "subtype_indirect"]),
        "figure3_contrast_rows": len(data["subtype_contrasts"]),
        "figure4_within_rows": len(data["attenuation_within"]),
        "figure4_contrast_rows": len(data["attenuation_contrasts"]),
    }
    expected = {
        "figure1_profile_rows": 15,
        "figure1_prevalence_rows": 6,
        "figure2_rows": 12,
        "figure3_indirect_rows": 9,
        "figure3_contrast_rows": 9,
        "figure4_within_rows": 9,
        "figure4_contrast_rows": 9,
    }
    failures = {k: {"expected": expected[k], "observed": v} for k, v in checks.items() if expected[k] != v}
    files = []
    for stem in [
        "Figure_1_subtype_landscape",
        "Figure_2_overall_mediation",
        "Figure_3_subtype_specific_mediation",
        "Figure_4_stress_si_attenuation",
    ]:
        for ext in [".svg", ".pdf", ".png", ".tiff"]:
            path = OUT_DIR / f"{stem}{ext}"
            files.append({"file": path.name, "exists": path.exists(), "bytes": path.stat().st_size if path.exists() else 0})
            if path.exists() and path.stat().st_size < 1000:
                failures[path.name] = {"expected": ">1000 bytes", "observed": path.stat().st_size}
        png = OUT_DIR / f"{stem}.png"
        if png.exists():
            img = Image.open(png)
            extrema = img.convert("L").getextrema()
            files.append({"file": png.name, "mode": img.mode, "size": img.size, "gray_extrema": extrema})
            if extrema[0] == extrema[1]:
                failures[png.name] = {"expected": "nonblank image", "observed": extrema}
    qa = {"checks": checks, "files": files, "failures": failures}
    (OUT_DIR / "generation_QA.json").write_text(json.dumps(qa, indent=2), encoding="utf-8")
    if failures:
        raise RuntimeError(json.dumps(failures, indent=2))
    return qa


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    data = extract_data()
    save_source_data(data)
    figure1(data)
    figure2(data)
    figure3(data)
    figure4(data)
    qa = validate_outputs(data)
    print(json.dumps({"status": "ok", "output_dir": str(OUT_DIR), "checks": qa["checks"]}, indent=2))


if __name__ == "__main__":
    main()
